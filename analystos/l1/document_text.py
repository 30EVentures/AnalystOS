"""L1 - read a document's real content as plain text, whatever shape it is.

Every other L1 extractor assumes the source *is* a table - a real
assumption that doesn't hold for a memo, a contract, a slide deck with
bullet points instead of a grid, or any document a person would actually
hand an analyst. ``extract_document_text`` makes no such assumption: it
reads whatever text content the format actually holds - paragraphs, bullet
points, table cells rendered as readable text - the same way regardless of
which of the five supported extensions it is. This is what makes "upload
anything, read it exactly as it is" real: one function, one text-in
contract, dispatched internally by extension.

Tables are rendered as plain, readable pipe-delimited text alongside
everything else, so a spreadsheet's numbers are still in the returned
text - but every table now goes through the exact same real,
type-aware table-parsing functions (``_raw_rows``, in each
``analystos.l1.extract*`` module) that ``analystos.l1.detect.extract_any``
already uses for the schema-driven path, not a separate, weaker
re-implementation that just joins raw cell text. This is one real
difference, not just a refactor: an Excel cell formatted as a percentage
stores its *fraction* (``0.571`` for what a person sees as "57.1%"), and
the old, separate flattening here showed that raw fraction verbatim -
``_raw_rows``'s percentage rescaling means this path now shows the same
human-correct "57.1" the schema path already did. See
specs/slice-44/spec.md.

DOCX footnotes are included too, each linked to the real body text it
sits near (``analystos.l1.footnotes.extract_footnotes_docx``) - see that
module for exactly which formats have a genuine structural footnote
concept to detect at all.
"""


import pdfplumber
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from analystos.l1 import extract, extract_docx, extract_pdf, extract_pptx, extract_xlsx
from analystos.l1.footnotes import extract_footnotes_docx
from analystos.l1.image_facts import extract_image_transcripts, render_image_blocks
from analystos.l1.pdf_columns import extract_page_text

SUPPORTED_EXTENSIONS = (".csv", ".docx", ".pdf", ".pptx", ".xlsx")


def _row_to_line(cells):
    return " | ".join((c or "").strip() for c in cells)


def _render_table(headers, numbered_rows):
    """One already-extracted table - ``(headers, numbered_rows)``, the
    exact shape ``_raw_rows`` returns - rendered as readable pipe-delimited
    text: the header row, then each data row in the same column order.
    The one place every format's table becomes text, so the rendering
    itself can't drift between formats the way five separate ad hoc
    flatteners could.
    """
    lines = [_row_to_line(headers)]
    for _row_num, raw in numbered_rows:
        lines.append(_row_to_line(raw.get(h, "") for h in headers))
    return "\n".join(lines)


def _text_from_csv(path, client=None):
    headers, rows = extract._raw_rows(path)
    return _render_table(headers, rows)


def _sheet_names(path):
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _text_from_xlsx(path, client=None):
    chunks = []
    for name in _sheet_names(path):
        try:
            headers, rows = extract_xlsx._raw_rows(path, sheet=name)
        except ValueError:
            continue  # an empty sheet - skip it, not fatal to the rest
        chunks.append(f"Sheet: {name}\n" + _render_table(headers, rows))
    return "\n\n".join(chunks)


def _text_from_docx(path, client=None):
    doc = Document(path)
    chunks = []
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if paragraphs:
        chunks.append("\n".join(paragraphs))

    for i in range(len(doc.tables)):
        try:
            headers, rows = extract_docx._raw_rows(path, table_index=i)
        except ValueError:
            continue  # an empty table - skip it, not fatal to the rest
        chunks.append(f"Table {i + 1}:\n" + _render_table(headers, rows))

    footnotes = extract_footnotes_docx(path)
    if footnotes:
        lines = ["Footnotes:"]
        for fn in footnotes:
            lines.append(f'[{fn["id"]}] (near: "{fn["marker_context"]}") {fn["footnote_text"]}')
        chunks.append("\n".join(lines))

    return "\n\n".join(chunks)


def _text_from_pptx(path, client=None):
    prs = Presentation(path)
    chunks = []
    table_i = 0  # a global index across the whole deck - matches how
    # extract_pptx._raw_rows numbers tables without a slide_index
    for slide_i, slide in enumerate(prs.slides, start=1):
        lines = [f"Slide {slide_i}:"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            elif shape.has_table:
                try:
                    headers, rows = extract_pptx._raw_rows(path, table_index=table_i)
                    lines.append(_render_table(headers, rows))
                except ValueError:
                    pass  # an empty table - skip it, not fatal to the rest
                table_i += 1
        if len(lines) > 1:
            chunks.append("\n".join(lines))
    return "\n\n".join(chunks)


def _text_from_pdf(path, client=None):
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page_i, page in enumerate(pdf.pages, start=1):
            lines = [f"Page {page_i}:"]
            text = extract_page_text(page)  # column-aware - see analystos.l1.pdf_columns
            if text:
                lines.append(text.strip())
            for table in page.extract_tables():
                try:
                    # The table this page already gave us goes through the same
                    # _raw_rows as the schema-driven path, without reopening the
                    # PDF and re-extracting every page per table (Slice 85).
                    headers, rows = extract_pdf._raw_rows(path, table=table)
                    lines.append(_render_table(headers, rows))
                except ValueError:
                    pass  # an empty table - skip it, not fatal to the rest
            if len(lines) > 1:
                chunks.append("\n".join(lines))

    # Embedded images (charts, scanned exhibits) transcribed via a real
    # vision call - a no-op, no-cost, no-client-required return of []
    # when the document has no embedded images at all. See
    # analystos.l1.image_facts.
    chunks.extend(render_image_blocks(extract_image_transcripts(path, client=client)))

    return "\n\n".join(chunks)


_EXTRACTORS = {
    ".csv": _text_from_csv,
    ".xlsx": _text_from_xlsx,
    ".docx": _text_from_docx,
    ".pptx": _text_from_pptx,
    ".pdf": _text_from_pdf,
}


def extract_document_text(path, client=None):
    """Return the real text content of the document at ``path``.

    Dispatches by extension - one of ``SUPPORTED_EXTENSIONS`` - to a
    format-specific reader, but every reader returns the same thing: plain
    text, paragraphs and rendered tables both included, every table
    parsed through the same real, type-aware logic
    ``analystos.l1.detect.extract_any`` uses (see the module docstring),
    never a separately re-implemented flattener. Raises ``ValueError`` for
    an unsupported extension or a document with no readable text at all.

    ``client`` is only ever used by the PDF path, to transcribe embedded
    images (``analystos.l1.image_facts``) - every other format ignores it.
    Defaults to ``None`` (a real client is resolved only if the document
    actually has an embedded image to transcribe; a plain PDF never
    touches it at all).
    """
    suffix = path.suffix.lower()
    if suffix not in _EXTRACTORS:
        raise ValueError(
            f"unsupported source file type {suffix!r}; "
            f"expected one of {sorted(SUPPORTED_EXTENSIONS)}"
        )
    text = _EXTRACTORS[suffix](path, client=client).strip()
    if not text:
        raise ValueError("document has no readable text")
    return text
