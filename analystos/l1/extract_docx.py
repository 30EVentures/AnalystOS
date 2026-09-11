"""L1 - pull a table out of a Word (.docx) document and check it against a schema.

Reads one table from the document - the first one unless ``table_index``
selects another. Row 1 is the header, row 2 the first data row - the same
convention as every other extractor, so a citation's row number means the
same thing regardless of source format.

A table in a .docx file is structured XML - a real table object with real
cells - not an inferred layout. Same reliability guarantee as CSV and Excel,
no review step needed.

Type-checking and number parsing are shared with every other format-specific
extractor - see ``analystos.l1.schema``.
"""

from docx import Document

from analystos.l1.schema import apply_schema


def _raw_rows(path, table_index=0):
    """Read table ``table_index`` from the .docx at ``path`` into
    ``(headers, numbered_raw_rows)`` - no schema applied yet. Shared by
    ``extract_table_docx`` and, for schema auto-detection,
    ``analystos.l1.detect.extract_any``.
    """
    doc = Document(path)
    if table_index >= len(doc.tables):
        raise ValueError(
            f"document has {len(doc.tables)} table(s); no table at index {table_index}"
        )
    table = doc.tables[table_index]

    if not table.rows:
        raise ValueError("table has no rows")

    headers = [c.text.strip() for c in table.rows[0].cells]

    numbered = []
    for row_num, row in enumerate(table.rows[1:], start=2):
        cells = row.cells
        if all(not c.text.strip() for c in cells):
            continue  # a blank row
        numbered.append((row_num, {h: c.text.strip() for h, c in zip(headers, cells)}))

    return headers, numbered


def all_tables(path):
    """Every table in the document, in document order, as
    ``[(headers, numbered_raw_rows), ...]`` - each extracted through the
    exact same ``_raw_rows`` the schema-driven path above already uses,
    not a separate re-implementation. A table with no rows is skipped
    (matches ``_raw_rows``'s own "empty table" refusal) rather than
    aborting the whole document. See specs/slice-44/spec.md - this is
    what ``analystos.l1.document_text`` now calls for a .docx's tables,
    instead of its own prior ad hoc cell-walking.
    """
    doc = Document(path)
    tables = []
    for i in range(len(doc.tables)):
        try:
            tables.append(_raw_rows(path, table_index=i))
        except ValueError:
            continue
    return tables


def extract_table_docx(path, schema, table_index=0):
    """Read table ``table_index`` from the .docx at ``path``; return its rows
    as a list of typed dicts.

    ``schema`` maps column name -> ``"number"`` or ``"text"``. Raises
    ``ValueError`` on an unknown schema type, a missing required column, a
    bad value, no table at ``table_index``, or an empty table.
    """
    headers, numbered = _raw_rows(path, table_index=table_index)
    return apply_schema(numbered, headers, schema)
