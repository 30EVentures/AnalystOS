"""L1 - pull a table out of an Excel (.xlsx) sheet and check it against a schema.

Reads the first sheet unless ``sheet`` names another one. Row 1 is the
header, row 2 the first data row - the same convention as the CSV extractor,
so a citation's row number means the same thing regardless of source format.

Excel stores a percentage-formatted cell as its fraction (``0.571`` for what
you see on screen as "57.1%"), not the display value. That's rescaled here so
a percentage read from a spreadsheet means the same thing as a percentage
AnalystOS computes itself (``57.1``, not ``0.571``) - otherwise a formatted
column would silently come out 100x too small.

Type-checking and number parsing are shared with every other format-specific
extractor - see ``analystos.l1.schema``.
"""

from openpyxl import load_workbook

from analystos.l1.schema import apply_schema


def _cell_to_raw(cell):
    """Turn one openpyxl cell into the string ``apply_schema`` expects."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if cell.number_format and "%" in cell.number_format:
            # round off binary float noise from the *100 (0.571 -> 57.1,
            # not 57.099999999999994) while keeping genuine precision
            value = round(value * 100, 6)
        return str(value)
    return str(value).strip()


def _raw_rows(path, sheet=None):
    """Read the sheet at ``path`` into ``(headers, numbered_raw_rows)`` - no
    schema applied yet. Shared by ``extract_table_xlsx`` and, for schema
    auto-detection, ``analystos.l1.detect.extract_any``.

    ``sheet`` selects a sheet by name; the workbook's first sheet is used
    otherwise. Raises ``ValueError`` on an empty sheet.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]

        rows_iter = ws.iter_rows()
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise ValueError("sheet has no rows")
        headers = [
            str(c.value).strip() if c.value is not None else "" for c in header_row
        ]

        numbered = []
        for row in rows_iter:
            if all(c.value is None for c in row):
                continue  # a blank row - common at the end of a real sheet
            row_num = row[0].row  # the actual 1-indexed sheet row, not a count
            numbered.append((row_num, {h: _cell_to_raw(c) for h, c in zip(headers, row)}))
    finally:
        wb.close()

    return headers, numbered


def extract_table_xlsx(path, schema, sheet=None):
    """Read the sheet at ``path``; return its rows as a list of typed dicts.

    ``schema`` maps column name -> ``"number"`` or ``"text"``. ``sheet``
    selects a sheet by name; the workbook's first sheet is used otherwise.
    Raises ``ValueError`` on an unknown schema type, a missing required
    column, a bad value, or an empty sheet.
    """
    headers, numbered = _raw_rows(path, sheet=sheet)
    return apply_schema(numbered, headers, schema)
