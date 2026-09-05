"""Tests for L1 extract_document_text - one per "Done when" in specs/slice-26/spec.md."""

import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet

from analystos.l1.document_text import extract_document_text


class ExtractDocumentTextTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_csv_returns_readable_text(self):
        path = self.tmp / "t.csv"
        path.write_text("period,revenue\nFY2024,4200000\n", encoding="utf-8")
        text = extract_document_text(path)
        self.assertIn("period", text)
        self.assertIn("4200000", text)

    def test_xlsx_returns_readable_text(self):
        path = self.tmp / "t.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["period", "revenue"])
        ws.append(["FY2024", 4200000])
        wb.save(path)
        text = extract_document_text(path)
        self.assertIn("Sheet: Data", text)
        self.assertIn("4200000", text)

    def test_docx_reads_prose_paragraphs_not_just_tables(self):  # the core ask
        path = self.tmp / "t.docx"
        doc = Document()
        doc.add_paragraph("Q3 was a strong quarter for the business overall.")
        doc.add_paragraph("Management expects continued momentum into Q4.")
        doc.save(path)
        text = extract_document_text(path)
        self.assertIn("Q3 was a strong quarter", text)
        self.assertIn("continued momentum", text)

    def test_docx_also_includes_table_content(self):
        path = self.tmp / "t.docx"
        doc = Document()
        doc.add_paragraph("Summary memo.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "period"
        table.cell(0, 1).text = "revenue"
        table.cell(1, 0).text = "FY2024"
        table.cell(1, 1).text = "4200000"
        doc.save(path)
        text = extract_document_text(path)
        self.assertIn("Summary memo.", text)
        self.assertIn("4200000", text)

    def test_pptx_reads_bullet_text_not_just_tables(self):
        path = self.tmp / "t.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, 100, 100)
        box.text_frame.text = "Revenue grew significantly this year."
        prs.save(path)
        text = extract_document_text(path)
        self.assertIn("Revenue grew significantly", text)

    def test_pdf_reads_prose_not_just_tables(self):
        path = self.tmp / "t.pdf"
        styles = getSampleStyleSheet()
        SimpleDocTemplate(str(path), pagesize=letter).build(
            [Paragraph("This is a real prose paragraph about the business.", styles["Normal"])]
        )
        text = extract_document_text(path)
        self.assertIn("real prose paragraph", text)

    def test_pdf_also_includes_table_content(self):
        path = self.tmp / "t.pdf"
        data = [["period", "revenue"], ["FY2024", "4200000"]]
        t = Table(data)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
        SimpleDocTemplate(str(path), pagesize=letter).build([t])
        text = extract_document_text(path)
        self.assertIn("4200000", text)

    def test_unsupported_extension_raises(self):
        path = self.tmp / "t.txt"
        path.write_text("hello", encoding="utf-8")
        with self.assertRaises(ValueError) as cm:
            extract_document_text(path)
        self.assertIn(".txt", str(cm.exception))

    def test_empty_document_raises(self):
        path = self.tmp / "t.docx"
        Document().save(path)
        with self.assertRaises(ValueError):
            extract_document_text(path)


if __name__ == "__main__":
    unittest.main()
