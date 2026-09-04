"""Tests for L1 extract_table_docx - one per "Done when" in specs/slice-17/spec.md."""

import tempfile
import unittest
from pathlib import Path

from docx import Document

from analystos.l1.extract_docx import extract_table_docx
from analystos.l2.answer import answer_lookup


def _make_docx(path, rows, extra_tables_before=0):
    """Write a minimal .docx with one table: rows[0] is the header.

    ``extra_tables_before`` adds that many throwaway 1-cell tables ahead of
    the real one, to test selecting a table by index.
    """
    doc = Document()
    for _ in range(extra_tables_before):
        doc.add_table(rows=1, cols=1)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row_vals in enumerate(rows):
        for c, val in enumerate(row_vals):
            table.cell(r, c).text = "" if val is None else str(val)
    doc.save(path)


class ExtractDocxTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_well_formed_table_returns_typed_rows(self):  # Done when #1
        path = self.tmp / "t.docx"
        _make_docx(path, [["period", "revenue"], ["FY2023", "3560000"], ["FY2024", "4200000"]])
        rows = extract_table_docx(path, {"period": "text", "revenue": "number"})
        self.assertEqual(
            rows, [{"period": "FY2023", "revenue": 3560000.0}, {"period": "FY2024", "revenue": 4200000.0}]
        )

    def test_spreadsheet_style_number_formats_are_accepted(self):  # shared with CSV/Excel
        path = self.tmp / "t.docx"
        _make_docx(path, [["period", "net_income"], ["FY2023", "$(4,368)"]])
        rows = extract_table_docx(path, {"period": "text", "net_income": "number"})
        self.assertEqual(rows[0]["net_income"], -4368.0)

    def test_blank_row_is_skipped_and_citation_row_stays_correct(self):  # Done when #2, cross-layer
        path = self.tmp / "t.docx"
        _make_docx(
            path,
            [["period", "revenue"], ["FY2023", "100"], [None, None], ["FY2024", "200"]],
        )
        rows = extract_table_docx(path, {"period": "text", "revenue": "number"})
        out = answer_lookup(rows, source="x" * 64, where=("period", "FY2024"), select="revenue")
        # table rows: 1=header, 2=FY2023, 3=blank, 4=FY2024 - not 3
        self.assertEqual(out["citation"]["row"], 4)

    def test_missing_required_column_raises(self):  # Done when #3
        path = self.tmp / "t.docx"
        _make_docx(path, [["period"], ["FY2024"]])
        with self.assertRaises(ValueError) as cm:
            extract_table_docx(path, {"period": "text", "revenue": "number"})
        self.assertIn("revenue", str(cm.exception))

    def test_table_index_selects_a_later_table(self):  # Done when #4
        path = self.tmp / "t.docx"
        _make_docx(
            path, [["period", "revenue"], ["FY2024", "4200000"]], extra_tables_before=1
        )
        rows = extract_table_docx(path, {"period": "text", "revenue": "number"}, table_index=1)
        self.assertEqual(rows, [{"period": "FY2024", "revenue": 4200000.0}])

    def test_out_of_range_table_index_raises(self):  # Done when #5
        path = self.tmp / "t.docx"
        _make_docx(path, [["period"], ["FY2024"]])
        with self.assertRaises(ValueError) as cm:
            extract_table_docx(path, {"period": "text"}, table_index=3)
        self.assertIn("1 table", str(cm.exception))

    def test_document_with_no_tables_raises(self):  # guard rail
        path = self.tmp / "t.docx"
        Document().save(path)
        with self.assertRaises(ValueError):
            extract_table_docx(path, {"period": "text"})


if __name__ == "__main__":
    unittest.main()
