"""Tests for L1 extract_table_pptx - one per "Done when" in specs/slice-18/spec.md."""

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from analystos.l1.extract_pptx import extract_table_pptx
from analystos.l2.answer import answer_lookup


def _make_pptx(path, rows, extra_slide_before=False):
    """Write a minimal .pptx with one table on one slide: rows[0] is the header."""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    if extra_slide_before:
        prs.slides.add_slide(blank)  # a slide with no table on it
    slide = prs.slides.add_slide(blank)
    shape = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(1), Inches(1), Inches(4), Inches(2)
    )
    table = shape.table
    for r, row_vals in enumerate(rows):
        for c, val in enumerate(row_vals):
            table.cell(r, c).text = "" if val is None else str(val)
    prs.save(path)


class ExtractPptxTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_well_formed_table_returns_typed_rows(self):  # Done when #1
        path = self.tmp / "t.pptx"
        _make_pptx(path, [["period", "revenue"], ["FY2023", "3560000"], ["FY2024", "4200000"]])
        rows = extract_table_pptx(path, {"period": "text", "revenue": "number"})
        self.assertEqual(
            rows, [{"period": "FY2023", "revenue": 3560000.0}, {"period": "FY2024", "revenue": 4200000.0}]
        )

    def test_blank_row_is_skipped_and_citation_row_stays_correct(self):  # Done when #2, cross-layer
        path = self.tmp / "t.pptx"
        _make_pptx(
            path, [["period", "revenue"], ["FY2023", "100"], [None, None], ["FY2024", "200"]]
        )
        rows = extract_table_pptx(path, {"period": "text", "revenue": "number"})
        out = answer_lookup(rows, source="x" * 64, where=("period", "FY2024"), select="revenue")
        self.assertEqual(out["citation"]["row"], 4)  # not 3

    def test_missing_required_column_raises(self):  # Done when #3
        path = self.tmp / "t.pptx"
        _make_pptx(path, [["period"], ["FY2024"]])
        with self.assertRaises(ValueError) as cm:
            extract_table_pptx(path, {"period": "text", "revenue": "number"})
        self.assertIn("revenue", str(cm.exception))

    def test_finds_the_first_table_anywhere_in_the_deck(self):  # Done when #4
        path = self.tmp / "t.pptx"
        _make_pptx(
            path, [["period", "revenue"], ["FY2024", "4200000"]], extra_slide_before=True
        )
        rows = extract_table_pptx(path, {"period": "text", "revenue": "number"})
        self.assertEqual(rows, [{"period": "FY2024", "revenue": 4200000.0}])

    def test_slide_index_narrows_the_search(self):  # Done when #5
        path = self.tmp / "t.pptx"
        _make_pptx(
            path, [["period", "revenue"], ["FY2024", "4200000"]], extra_slide_before=True
        )
        # the table is on slide 1 (0-based) - slide 0 has none
        with self.assertRaises(ValueError) as cm:
            extract_table_pptx(path, {"period": "text"}, slide_index=0)
        self.assertIn("slide 0", str(cm.exception))
        rows = extract_table_pptx(path, {"period": "text", "revenue": "number"}, slide_index=1)
        self.assertEqual(rows, [{"period": "FY2024", "revenue": 4200000.0}])

    def test_deck_with_no_tables_raises(self):  # guard rail
        path = self.tmp / "t.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(path)
        with self.assertRaises(ValueError):
            extract_table_pptx(path, {"period": "text"})


if __name__ == "__main__":
    unittest.main()
