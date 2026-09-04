"""Tests for L1 extract_table_xlsx - one per "Done when" in specs/slice-16/spec.md."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from analystos.l1.extract_xlsx import extract_table_xlsx
from analystos.l2.answer import answer_lookup


def _make_xlsx(path, rows, sheet_name=None, percent_cells=()):
    """Write a minimal .xlsx: rows[0] is the header, the rest are data.

    ``percent_cells`` is a set of (row_index_in_rows, col_index) to format as
    a percentage - openpyxl then stores the value as its fraction, matching
    what a real "Format Cells > Percentage" column looks like.
    """
    wb = Workbook()
    ws = wb.active
    if sheet_name:
        ws.title = sheet_name
    for r in rows:
        ws.append(r)
    for r, c in percent_cells:
        ws.cell(row=r + 1, column=c + 1).number_format = "0.0%"
    wb.save(path)
    wb.close()


class ExtractXlsxTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_well_formed_sheet_returns_typed_rows(self):  # Done when #1
        path = self.tmp / "t.xlsx"
        _make_xlsx(path, [["period", "revenue"], ["FY2023", 3560000], ["FY2024", 4200000]])
        rows = extract_table_xlsx(path, {"period": "text", "revenue": "number"})
        self.assertEqual(
            rows, [{"period": "FY2023", "revenue": 3560000.0}, {"period": "FY2024", "revenue": 4200000.0}]
        )

    def test_percentage_formatted_cell_is_rescaled(self):  # Done when #2
        path = self.tmp / "t.xlsx"
        # a cell showing "57.1%" in Excel is stored internally as 0.571
        _make_xlsx(path, [["period", "margin"], ["FY2024", 0.571]], percent_cells={(1, 1)})
        rows = extract_table_xlsx(path, {"period": "text", "margin": "number"})
        self.assertEqual(rows[0]["margin"], 57.1)

    def test_blank_row_is_skipped_and_later_row_numbers_stay_accurate(self):  # Done when #3
        path = self.tmp / "t.xlsx"
        _make_xlsx(
            path,
            [["period", "revenue"], ["FY2023", 100], [None, None], ["FY2024", "bad"]],
        )
        with self.assertRaises(ValueError) as cm:
            extract_table_xlsx(path, {"period": "text", "revenue": "number"})
        # FY2024 is sheet row 4 (header=1, FY2023=2, blank=3, FY2024=4) -
        # not row 3, which is where it would land if the blank row shifted
        # a position-based count instead of using the sheet's real row number.
        self.assertIn("row 4", str(cm.exception))

    def test_missing_required_column_raises(self):  # Done when #4
        path = self.tmp / "t.xlsx"
        _make_xlsx(path, [["period"], ["FY2024"]])
        with self.assertRaises(ValueError) as cm:
            extract_table_xlsx(path, {"period": "text", "revenue": "number"})
        self.assertIn("revenue", str(cm.exception))

    def test_named_sheet_is_selected(self):  # Done when #5
        path = self.tmp / "t.xlsx"
        wb = Workbook()
        wb.active.title = "Cover"
        wb.active.append(["ignore", "me"])
        data = wb.create_sheet("Income Statement")
        data.append(["period", "revenue"])
        data.append(["FY2024", 4200000])
        wb.save(path)
        wb.close()

        rows = extract_table_xlsx(
            path, {"period": "text", "revenue": "number"}, sheet="Income Statement"
        )
        self.assertEqual(rows, [{"period": "FY2024", "revenue": 4200000.0}])

    def test_citation_row_is_correct_in_l2_after_a_skipped_row(self):  # Done when #3, cross-layer
        # The bug this guards: L1 knows the real sheet row of every row it
        # returns, but a citation is built later, in L2, from the row's
        # *position* in the returned list. If a blank row was skipped during
        # extraction, position and sheet row diverge - position 1 is sheet
        # row 4, not row 3 - and a naive "position + 2" citation would point
        # at the wrong cell. This proves the real row number survives the
        # handoff from L1 to L2 instead.
        path = self.tmp / "t.xlsx"
        _make_xlsx(path, [["period", "revenue"], ["FY2023", 100], [None, None], ["FY2024", 200]])
        rows = extract_table_xlsx(path, {"period": "text", "revenue": "number"})
        out = answer_lookup(rows, source="x" * 64, where=("period", "FY2024"), select="revenue")
        self.assertEqual(out["citation"]["row"], 4)  # the real sheet row, not 3

    def test_empty_sheet_raises(self):  # guard rail
        path = self.tmp / "t.xlsx"
        Workbook().save(path)
        with self.assertRaises(ValueError):
            extract_table_xlsx(path, {"period": "text"})


if __name__ == "__main__":
    unittest.main()
