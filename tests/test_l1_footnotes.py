"""Tests for real, structural footnote-marker-to-content linking -
specs/slice-44/spec.md. Only .docx has a genuine structural footnote
relationship (see analystos.l1.footnotes' module docstring for why every
other supported format doesn't) - these tests build a real OOXML footnote
by hand, since neither python-docx's public API nor any fixture file in
this repo has one.
"""

import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from analystos.l1.footnotes import extract_footnotes_docx

_FOOTNOTES_XML = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
<w:footnote w:id="2"><w:p><w:r><w:t xml:space="preserve">Unaudited; subject to year-end adjustment.</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""

_CONTENT_TYPE_OVERRIDE = (
    '<Override PartName="/word/footnotes.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
)
_FOOTNOTES_RELATIONSHIP = (
    '<Relationship Id="rIdFootnotes1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
    'Target="footnotes.xml"/>'
)


def _add_footnote_reference(paragraph, footnote_id):
    """Inject a real ``<w:footnoteReference w:id="...">`` run into
    ``paragraph`` - python-docx's public API has no method for this, so
    this builds the OOXML directly the way python-docx's own low-level
    customization recipes do."""
    run_el = etree.SubElement(paragraph._p, qn("w:r"))
    ref_el = etree.SubElement(run_el, qn("w:footnoteReference"))
    ref_el.set(qn("w:id"), str(footnote_id))


def _build_docx_with_footnote(path):
    """A real, minimal .docx with a genuine OOXML footnote relationship:
    a body paragraph carrying a real footnoteReference, a footnotes.xml
    part with the matching footnote content, the content-type override,
    and the part relationship - the same structure Word itself produces,
    not a simplified stand-in.
    """
    doc = Document()
    doc.add_paragraph("Revenue was $10 million, up from the prior year")
    _add_footnote_reference(doc.paragraphs[-1], footnote_id=2)
    doc.add_paragraph("A second paragraph with no footnote at all")
    doc.save(path)

    with zipfile.ZipFile(path) as zin:
        items = {name: zin.read(name) for name in zin.namelist()}

    content_types = items["[Content_Types].xml"].decode("utf-8")
    content_types = content_types.replace("</Types>", _CONTENT_TYPE_OVERRIDE + "</Types>")
    items["[Content_Types].xml"] = content_types.encode("utf-8")

    rels = items["word/_rels/document.xml.rels"].decode("utf-8")
    rels = rels.replace("</Relationships>", _FOOTNOTES_RELATIONSHIP + "</Relationships>")
    items["word/_rels/document.xml.rels"] = rels.encode("utf-8")

    items["word/footnotes.xml"] = _FOOTNOTES_XML

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items.items():
            zout.writestr(name, data)


class DocxFootnoteLinkingTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "with_footnote.docx"
        _build_docx_with_footnote(self.path)

    def tearDown(self):
        self._tmp.cleanup()

    def test_the_footnote_marker_is_linked_to_its_real_content(self):  # Done when (Task 2)
        linked = extract_footnotes_docx(self.path)
        self.assertEqual(len(linked), 1)
        entry = linked[0]
        self.assertEqual(entry["id"], "2")
        self.assertEqual(entry["footnote_text"], "Unaudited; subject to year-end adjustment.")

    def test_the_marker_context_identifies_where_in_the_body_it_sits(self):
        # Retrievable as associated with its marker, not just floating
        # prose - the footnote text alone proves nothing about *where* it
        # belongs without this.
        linked = extract_footnotes_docx(self.path)
        self.assertIn("Revenue was $10 million", linked[0]["marker_context"])
        self.assertNotIn("second paragraph", linked[0]["marker_context"])

    def test_a_document_with_no_footnotes_returns_an_empty_list(self):
        plain_path = Path(self._tmp.name) / "plain.docx"
        Document().save(plain_path)  # python-docx never adds a footnotes part on its own
        doc = Document(plain_path)
        doc.add_paragraph("Nothing footnoted here at all.")
        doc.save(plain_path)
        self.assertEqual(extract_footnotes_docx(plain_path), [])

    def test_separator_and_continuation_separator_entries_are_never_returned_as_content(self):
        # Every real Word document carries these two structurally (id -1
        # and 0) - they are page-break markup, never real footnote text,
        # and must never be mistaken for one.
        linked = extract_footnotes_docx(self.path)
        ids = {entry["id"] for entry in linked}
        self.assertNotIn("-1", ids)
        self.assertNotIn("0", ids)


if __name__ == "__main__":
    unittest.main()
