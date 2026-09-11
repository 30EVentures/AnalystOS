"""Integration test for GAAP/non-GAAP (parallel-metric) detection -
specs/slice-46/spec.md, Task 1. Proves the real seam: a real reconciliation
table, extracted through L1's unified table-parsing path (Slice 44), feeds
L2's gaap_status tagging and verification (Slice 45) - not two isolated
unit tests each asserting the other's contract on faith.

Runs the narrated-default extraction path specifically:
analystos.l1.document_text.extract_document_text (which now calls the
same _raw_rows every format module's schema-driven path already uses) ->
analystos.l2.analyze.analyze_document. The old schema/template path
(analystos.l1.detect.extract_any) never touches prose or a reconciliation
table's own row/column framing the way this path does, so it isn't what's
being exercised here.
"""

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from analystos.l1 import extract_docx
from analystos.l1.document_text import extract_document_text
from analystos.l2.analyze import analyze_document

TITLE = "Q3 2026 reconciliation"


def _build_reconciliation_docx(path):
    doc = Document()
    doc.add_paragraph("Reconciliation of GAAP to non-GAAP net income:")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Line item", "Amount"
    table.cell(1, 0).text, table.cell(1, 1).text = "Net income (GAAP)", "$19,600,000"
    table.cell(2, 0).text, table.cell(2, 1).text = "Net income (Non-GAAP)", "$24,100,000"
    doc.save(path)


def _fake_client(segments):
    tool_use_block = SimpleNamespace(type="tool_use", input={"segments": segments})
    response = SimpleNamespace(content=[tool_use_block], stop_reason="tool_use")
    return SimpleNamespace(messages=SimpleNamespace(create=lambda **kwargs: response))


def _plausible_model_response():
    """What a real model reading the extracted table text would plausibly
    return - both headline figures tagged by basis, plus the reconciling
    gap as an ordinary "difference" computation."""
    return [
        {
            "type": "quote", "gaap_status": "gaap", "display": "stat",
            "label": "Net income (GAAP)", "exact_text": "$19,600,000",
            "has_value": True, "value": 19600000.0,
            "sentence": "GAAP net income was {value}.", "format": "usd",
        },
        {
            "type": "quote", "gaap_status": "non_gaap", "display": "stat",
            "label": "Net income (Non-GAAP)", "exact_text": "$24,100,000",
            "has_value": True, "value": 24100000.0,
            "sentence": "Non-GAAP net income was {value}.", "format": "usd",
        },
        {
            "type": "computed", "gaap_status": "n/a", "display": "inline",
            "label": "GAAP/non-GAAP reconciling gap", "operation": "difference",
            "operands": [
                {"exact_text": "$24,100,000", "value": 24100000.0},
                {"exact_text": "$19,600,000", "value": 19600000.0},
            ],
            "has_total": False, "total_exact_text": "", "total_value": 0,
            "result": 4500000.0, "sentence": "The reconciling gap was {value}.",
            "format": "usd",
        },
    ]


class GaapReconciliationTableIntegrationTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "reconciliation.docx"
        _build_reconciliation_docx(self.path)

    def tearDown(self):
        self._tmp.cleanup()

    def test_the_table_is_extracted_through_the_real_unified_table_path(self):
        # Confirms which extraction path this runs on, directly - not
        # asserted, proven: the real _raw_rows every schema-driven table
        # already goes through is what actually parsed this table.
        with patch.object(extract_docx, "_raw_rows", wraps=extract_docx._raw_rows) as spy:
            text = extract_document_text(self.path)
        self.assertGreaterEqual(spy.call_count, 1)
        self.assertIn("Line item | Amount", text)
        self.assertIn("Net income (GAAP) | $19,600,000", text)
        self.assertIn("Net income (Non-GAAP) | $24,100,000", text)

    def test_both_figures_verify_against_the_real_table_and_are_linked_by_basis(self):
        document_text = extract_document_text(self.path)
        client = _fake_client(_plausible_model_response())

        out = analyze_document(document_text, TITLE, client=client)

        self.assertEqual(len(out), 3)
        gaap = next(s for s in out if s.get("gaap_status") == "gaap")
        non_gaap = next(s for s in out if s.get("gaap_status") == "non_gaap")
        gap = next(s for s in out if s["type"] == "computed")

        # not extracted as two unrelated numbers: both are real, verified
        # figures from the same table, correctly tagged by basis...
        self.assertEqual(gaap["value"], 19600000.0)
        self.assertEqual(non_gaap["value"], 24100000.0)
        # ...and linked - the reconciling gap is itself a real, verified
        # fact computed from exactly these same two numbers.
        self.assertEqual(gap["value"], 4500000.0)
        self.assertEqual(gap["gaap_status"], "n/a")

    def test_a_fabricated_figure_not_actually_in_the_table_fails_verification(self):
        # Defense in depth: even if a model hallucinated a third "basis"
        # figure the table never stated, citation-matching against the
        # real extracted text refuses it - never a fabricated link.
        document_text = extract_document_text(self.path)
        bad_segments = _plausible_model_response() + [{
            "type": "quote", "gaap_status": "non_gaap", "display": "stat",
            "label": "Adjusted EBITDA", "exact_text": "$31,000,000",
            "has_value": True, "value": 31000000.0,
            "sentence": "Adjusted EBITDA was {value}.", "format": "usd",
        }]
        out = analyze_document(document_text, TITLE, client=_fake_client(bad_segments))
        self.assertEqual(len(out), 3)  # the fabricated 4th segment was dropped
        self.assertFalse(any(s.get("value") == 31000000.0 for s in out))


if __name__ == "__main__":
    unittest.main()
