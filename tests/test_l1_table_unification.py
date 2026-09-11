"""Proves the narrated-default path (analystos.l1.document_text) and the
schema-driven path (analystos.l1.detect.extract_any) now parse a table
through the exact same real, type-aware function - one per format, five
total. Before this fix, document_text.py maintained its own separate,
weaker flattening that never called ``_raw_rows`` at all; these tests
fail under that old behavior (the spy would record zero calls) and pass
now. See specs/slice-44/spec.md.
"""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from analystos.l1 import extract, extract_docx, extract_pdf, extract_pptx, extract_xlsx
from analystos.l1.detect import extract_any
from analystos.l1.document_text import extract_document_text

_STYLES = getSampleStyleSheet()


def _reportlab_table(path, rows):
    data = [[str(v) for v in row] for row in rows]
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(str(path), pagesize=letter).build([t])


class TableUnificationTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_csv_uses_the_same_raw_rows_as_the_schema_driven_path(self):  # Done when (Task 1)
        path = self.tmp / "t.csv"
        path.write_text("period,revenue\nFY2024,4200000\nFY2025,5100000\n", encoding="utf-8")

        with patch.object(extract, "_raw_rows", wraps=extract._raw_rows) as spy:
            text = extract_document_text(path)
        self.assertGreaterEqual(spy.call_count, 1)  # would be 0 under the old separate flattener

        schema, rows = extract_any(path)
        self.assertEqual(schema, {"period": "text", "revenue": "number"})
        self.assertEqual(rows[0]["revenue"], 4200000.0)
        # the narrated path's text carries the same real values the typed
        # path independently parsed - not a coincidence, the same function
        self.assertIn("period | revenue", text)
        self.assertIn("FY2024 | 4200000", text)

    def test_docx_uses_the_same_raw_rows_as_the_schema_driven_path(self):  # Done when (Task 1)
        path = self.tmp / "t.docx"
        doc = Document()
        doc.add_paragraph("Summary memo.")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "period", "revenue"
        table.cell(1, 0).text, table.cell(1, 1).text = "FY2024", "4200000"
        table.cell(2, 0).text, table.cell(2, 1).text = "", ""  # a blank row
        doc.save(path)

        with patch.object(extract_docx, "_raw_rows", wraps=extract_docx._raw_rows) as spy:
            text = extract_document_text(path)
        self.assertGreaterEqual(spy.call_count, 1)  # would be 0 under the old separate flattener

        schema, rows = extract_any(path)
        self.assertEqual(rows[0]["revenue"], 4200000.0)
        self.assertIn("period | revenue", text)
        self.assertIn("FY2024 | 4200000", text)
        # the blank row: _raw_rows skips it, the old ad hoc cell-walk
        # never did - a real structural fix, not just a refactor
        self.assertNotIn(" |  ", text)

    def test_xlsx_uses_the_same_raw_rows_as_the_schema_driven_path(self):  # Done when (Task 1)
        path = self.tmp / "t.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["period", "gross_margin_pct"])
        ws.append(["FY2024", 0.571])
        ws["B2"].number_format = "0.0%"  # stored as a fraction, displays as 57.1%
        wb.save(path)

        with patch.object(extract_xlsx, "_raw_rows", wraps=extract_xlsx._raw_rows) as spy:
            text = extract_document_text(path)
        self.assertGreaterEqual(spy.call_count, 1)  # would be 0 under the old separate flattener

        schema, rows = extract_any(path)
        self.assertEqual(rows[0]["gross_margin_pct"], 57.1)
        # the real bug this fixes: the old flattener showed openpyxl's raw
        # stored value (0.571) verbatim - a person reading the sheet sees
        # "57.1%", not "0.571". The narrated path must show the same
        # human-correct value the schema path already did.
        self.assertIn("57.1", text)
        self.assertNotIn("0.571", text)

    def test_pptx_uses_the_same_raw_rows_as_the_schema_driven_path(self):  # Done when (Task 1)
        path = self.tmp / "t.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_table(3, 2, 0, 0, 2000000, 1000000)
        table = shape.table
        table.cell(0, 0).text, table.cell(0, 1).text = "period", "revenue"
        table.cell(1, 0).text, table.cell(1, 1).text = "FY2024", "4200000"
        table.cell(2, 0).text, table.cell(2, 1).text = "", ""  # a blank row
        prs.save(path)

        with patch.object(extract_pptx, "_raw_rows", wraps=extract_pptx._raw_rows) as spy:
            text = extract_document_text(path)
        self.assertGreaterEqual(spy.call_count, 1)  # would be 0 under the old separate flattener

        schema, rows = extract_any(path)
        self.assertEqual(rows[0]["revenue"], 4200000.0)
        self.assertIn("period | revenue", text)
        self.assertIn("FY2024 | 4200000", text)
        self.assertNotIn(" |  ", text)  # the blank row is skipped, matching _raw_rows

    def test_pdf_uses_the_same_raw_rows_as_the_schema_driven_path(self):  # Done when (Task 1)
        path = self.tmp / "t.pdf"
        _reportlab_table(path, [["period", "revenue"], ["FY2024", "4200000"]])

        with patch.object(extract_pdf, "_raw_rows", wraps=extract_pdf._raw_rows) as spy:
            text = extract_document_text(path)
        self.assertGreaterEqual(spy.call_count, 1)  # would be 0 under the old separate flattener

        schema, rows = extract_any(path)
        self.assertEqual(rows[0]["revenue"], 4200000.0)
        self.assertIn("period | revenue", text)
        self.assertIn("FY2024 | 4200000", text)


if __name__ == "__main__":
    unittest.main()
